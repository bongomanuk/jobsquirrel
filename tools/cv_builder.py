from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date
import json
import os
import re

OUT = os.environ.get("CV_OUT", os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "documents", "applications"))
os.makedirs(OUT, exist_ok=True)

_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG_DIR = os.path.join(_REPO, "config")
HEADSHOT = os.environ.get("CV_HEADSHOT", os.path.join(_REPO, "documents", "assets", "headshot.png"))


def _config_path(name):
    """Prefer config/<name>; fall back to the matching config/<stem>.example<suffix>."""
    path = os.path.join(CONFIG_DIR, name)
    if os.path.exists(path):
        return path
    stem, ext = os.path.splitext(name)
    return os.path.join(CONFIG_DIR, stem + ".example" + ext)


def _read_json(name):
    """Load a JSON config file, or return None if it is missing or unreadable."""
    try:
        with open(_config_path(name), encoding="utf-8") as fh:
            return json.load(fh)
    except (FileNotFoundError, json.JSONDecodeError):
        return None

def dedash(text):
    """Never allow em dashes or ' -- ' in rendered documents; use commas."""
    if not isinstance(text, str):
        return text
    text = text.replace(' \u2014 ', ', ').replace('\u2014', ', ')
    text = text.replace(' -- ', ', ')
    text = re.sub(r'\s+,', ',', text)
    text = re.sub(r',\s*,', ',', text)
    return text


DARK = RGBColor(0x1a, 0x1a, 0x1a)
GREY = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x0b, 0x3d, 0x66)

def _read_profile():
    """Identity and eligibility answers parsed from config/profile.md (or the example)."""
    vals = {}
    try:
        with open(_config_path("profile.md"), encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if line.startswith("- ") and ":" in line:
                    key, value = line[2:].split(":", 1)
                    value = value.strip()
                    if value:
                        vals[key.strip().lower()] = value
    except FileNotFoundError:
        pass
    return vals


_ANSWERS = _read_profile()
_CV = _read_json("cv.json") or {}


def _identity(env_key, answer_key, default=""):
    return os.environ.get(env_key) or _ANSWERS.get(answer_key, "") or default


NAME = _identity("CV_NAME", "full name", _CV.get("name", "YOUR NAME"))
_EMAIL = _identity("CV_EMAIL", "email")
_PHONE = _identity("CV_PHONE", "phone")
_LOCATION = _identity("CV_LOCATION", "location (city, country)")
_LINKEDIN = _identity("CV_LINKEDIN", "linkedin")
_WEBSITE = _identity("CV_WEBSITE", "github / portfolio / website")
_OTHER_SITES = [s.strip() for s in _identity("CV_OTHER_SITES", "other websites (comma-separated)").split(",") if s.strip()]

CVD_NAME = os.environ.get("CV_DISPLAY_NAME") or NAME

CONTACT_LINE = os.environ.get("CV_CONTACT") or "  |  ".join(x for x in (_LOCATION, _PHONE, _EMAIL) if x)
LINKS_LINE = os.environ.get("CV_LINKS") or "  |  ".join(x for x in ([_WEBSITE, _LINKEDIN] + _OTHER_SITES) if x)
MOTIVATION_CONTACT = "  |  ".join(x for x in (_PHONE, _EMAIL, _LINKEDIN) if x)


def _filename_prefix():
    """Filename prefix: 'First_Last' from the profile full name, or $CV_PREFIX, or 'CV'."""
    override = os.environ.get("CV_PREFIX")
    if override:
        return override
    parts = _ANSWERS.get("full name", "").split()
    if len(parts) >= 2:
        return parts[0] + "_" + parts[-1]
    if parts:
        return parts[0]
    return "CV"


CV_PREFIX = _filename_prefix()

def _load_experience():
    """Experience from config/cv.json, shaped for the renderers (org/dates/core/extra)."""
    out = []
    for e in _CV.get("experience", []) or []:
        org = (e.get("employer", "") or "").strip()
        title = (e.get("title", "") or "").strip()
        dates = (e.get("dates", "") or "").strip()
        location = (e.get("location", "") or "").strip()
        out.append({
            "org": (org + "  |  " + title) if title else org,
            "dates": (dates + "  |  " + location) if location else dates,
            "core": list(e.get("bullets", []) or []),
            "extra": dict(e.get("extras", {}) or {}),
        })
    return out


EXPERIENCE = _load_experience()
PROJECTS = list(_CV.get("projects", []) or [])
EDUCATION = list(_CV.get("education", []) or [])

def _is_contract(role):
    return str(role.get("engagement", "fulltime")).lower() in ("contract", "contracting", "consulting", "consultant")


def _rate(role):
    return role.get("rate") or _ANSWERS.get("contract rate (ex vat)", "")


def _notice_clause(notice):
    if not notice:
        return ""
    if notice.strip().lower() in ("none", "n/a", "na", "0"):
        return " with no notice period"
    return " with a notice period of " + notice


def availability_text(role):
    parts = []
    earliest = _ANSWERS.get("earliest start", "")
    notice = _ANSWERS.get("notice period", "")
    if earliest:
        parts.append("Available: " + earliest)
    if notice:
        parts.append("Notice period: " + notice)
    if _is_contract(role):
        rate = _rate(role)
        if rate:
            parts.append("Rate: " + rate)
        holidays = _ANSWERS.get("planned holidays", "")
        if holidays:
            parts.append("Planned holidays: " + holidays)
    right_to_work = _ANSWERS.get("right to work", "")
    if right_to_work:
        parts.append("Work authorisation: " + right_to_work)
    return "  |  ".join(parts)


def practical_text(role):
    """The 'practical details' paragraph. A day rate only belongs in a contract application."""
    earliest = _ANSWERS.get("earliest start", "")
    notice = _ANSWERS.get("notice period", "")
    location = _ANSWERS.get("location (city, country)", "")
    right_to_work = _ANSWERS.get("right to work", "")
    parts = ["On the practical details:"]
    available = ("I am available " + earliest) if earliest else "I am"
    parts.append((available + _notice_clause(notice)).strip() + ".")
    if _is_contract(role):
        rate = _rate(role)
        if rate:
            parts.append("My rate is " + rate + ".")
        holidays = _ANSWERS.get("planned holidays", "")
        if holidays:
            parts.append("Planned holidays: " + holidays + ".")
    if location:
        parts.append("I am based in " + location + ".")
    if right_to_work:
        parts.append("Work authorisation: " + right_to_work + ".")
    parts.append("My CV is attached separately with the full picture.")
    return " ".join(parts)


def _letter_date():
    override = os.environ.get("CV_DATE")
    if override:
        return override
    location = _ANSWERS.get("location (city, country)", "")
    city = location.split(",")[0].strip() if location else ""
    when = date.today().strftime("%-d %B %Y")
    return (city + ", " + when) if city else when


LETTER_DATE = _letter_date()


# ---------------- doc helpers ----------------
def new_doc(font='Calibri'):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = font
    st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), font)
    for s in doc.sections:
        s.page_width = Cm(21.0); s.page_height = Cm(29.7)   # A4
        s.top_margin = Cm(1.2); s.bottom_margin = Cm(1.0)
        s.left_margin = Cm(1.4); s.right_margin = Cm(1.4)
    return doc


def heading(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(2); pf.line_spacing = 1.0
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = ACCENT
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6'); bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), '0B3D66')
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def body(doc, text, size=9.5, color=DARK, space_after=3, bold=False):
    text = dedash(text)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold
    return p


def bullet(doc, text, size=9.5):
    text = dedash(text)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5); pf.first_line_indent = Cm(-0.35); pf.space_after = Pt(1.5); pf.space_before = Pt(0); pf.line_spacing = 1.0
    r0 = p.add_run("\u2022  "); r0.font.size = Pt(size); r0.font.color.rgb = DARK
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = DARK
    return p


def header(doc, headline, rules=True):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(NAME.upper()); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = DARK
    if headline:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(headline); r.font.size = Pt(10.5); r.font.color.rgb = ACCENT
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(CONTACT_LINE); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(LINKS_LINE); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    if rules:
        hr = doc.add_paragraph(); hr.paragraph_format.space_after = Pt(4); hr.paragraph_format.space_before = Pt(0)
        pPr = hr._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '8'); bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '0B3D66')
        pbdr.append(bottom); pPr.append(pbdr)


# ---------------- renderers ----------------
# Two-column CV matching the reference CTO template: a full-bleed navy sidebar
# (left) with a circular headshot, and a serif main column.

SIDEBAR_FILL = "173D56"
NAVY = RGBColor(0x17, 0x3D, 0x56)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def _load_sidebar():
    """Sidebar content from config/cv.json.

    A `sidebar` block is optional; older configs keep education, certifications
    and languages at the top level, so use those as the fallback.
    """
    sb = _CV.get("sidebar", {}) or {}
    return {
        "education": list(sb.get("education") or _CV.get("education", []) or []),
        "certificates": list(sb.get("certificates") or _CV.get("certifications", []) or []),
        "leadership": list(sb.get("leadership", []) or []),
        "languages": list(sb.get("languages") or _CV.get("languages", []) or []),
        "strengths": list(sb.get("strengths", []) or []),
        "ai_ml": list(sb.get("ai_ml", []) or []),
    }


SIDEBAR = _load_sidebar()


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'none'); borders.append(e)
    tblPr.append(borders)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _cell_margins(cell, top=90, left=150, bottom=90, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + tag); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcPr.append(mar)


def _row_at_least(row, cm_h):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(int(cm_h * 567))); h.set(qn('w:hRule'), 'atLeast')
    trPr.append(h)


def _tracking(run, tw):
    rPr = run._element.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(tw)); rPr.append(sp)


def s_head(cell, text):
    p = cell.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(1.5); pf.line_spacing = 1.0
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = WHITE
    _tracking(r, 30)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'FFFFFF')
    pbdr.append(b); pPr.append(pbdr)


def s_text(cell, text, size=8.5, bullet=False, bold=False):
    text = dedash(text)
    p = cell.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(1.0); pf.line_spacing = 1.0
    if bullet:
        pf.left_indent = Cm(0.3); pf.first_line_indent = Cm(-0.3)
        r0 = p.add_run("\u2022  "); r0.font.size = Pt(size); r0.font.color.rgb = WHITE
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = WHITE; r.bold = bold
    return p


def m_head(cell, text):
    p = cell.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(7); pf.space_after = Pt(2); pf.line_spacing = 1.0
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NAVY
    _tracking(r, 20)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:space'), '2'); b.set(qn('w:color'), '173D56')
    pbdr.append(b); pPr.append(pbdr)


def m_body(cell, text, size=9.5, color=DARK, italic=False, bold=False, space_after=3):
    text = dedash(text)
    p = cell.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(space_after); pf.line_spacing = 1.0
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color; r.italic = italic; r.bold = bold
    return p


def m_bullet(cell, text, size=9.5):
    text = dedash(text)
    p = cell.add_paragraph(); pf = p.paragraph_format
    pf.left_indent = Cm(0.55); pf.first_line_indent = Cm(-0.4)
    pf.space_after = Pt(1.0); pf.space_before = Pt(0); pf.line_spacing = 1.0
    r0 = p.add_run("\u2022   "); r0.font.size = Pt(size); r0.font.color.rgb = NAVY
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = DARK
    return p


def m_entry(cell, org, dates, title, bullets, extras, width_cm, groups=None):
    from docx.enum.text import WD_TAB_ALIGNMENT
    org, dates = dedash(org), dedash(dates)
    emp, _, ttl = org.partition("  |  ")
    period = dates.split("  |  ")[0]
    p = cell.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(4); pf.space_after = Pt(0); pf.line_spacing = 1.0
    pf.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(emp.strip()); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = NAVY
    r2 = p.add_run("\t" + period.strip()); r2.font.size = Pt(9); r2.font.color.rgb = DARK
    role_title = ttl.strip() or title
    if role_title:
        p = cell.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(role_title); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
    for b in bullets:
        m_bullet(cell, b, size=9.0)
    for g in (groups or []):
        label = g.get("label", "")
        if label:
            p = cell.add_paragraph(); pf = p.paragraph_format
            pf.space_before = Pt(3); pf.space_after = Pt(0.5); pf.line_spacing = 1.0; pf.left_indent = Cm(0.1)
            r = p.add_run(label); r.bold = True; r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
        for b in g.get("bullets", []):
            m_bullet(cell, b, size=9.0)
    for b in extras:
        m_bullet(cell, b)


def add_page_bar(doc, path, w_cm, h_cm):
    from docx.oxml.ns import qn
    sec = doc.sections[0]
    try: sec.header_distance = Cm(0)
    except Exception: pass
    hdr = sec.header
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]
    if p.runs: p.runs[0].text = ""
    run = p.add_run()
    run.add_picture(path, width=Cm(w_cm), height=Cm(h_cm))
    drawing = run._element.find(qn('w:drawing'))
    inline = drawing.find(qn('wp:inline')) if drawing is not None else None
    if inline is None:
        return
    anchor = OxmlElement('wp:anchor')
    for k, v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0',
                 'relativeHeight':'1','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():
        anchor.set(k, v)
    sp = OxmlElement('wp:simplePos'); sp.set('x','0'); sp.set('y','0'); anchor.append(sp)
    for tag in ('wp:positionH','wp:positionV'):
        pos = OxmlElement(tag); pos.set('relativeFrom','page')
        off = OxmlElement('wp:posOffset'); off.text = '0'; pos.append(off); anchor.append(pos)
    for tag in ('wp:extent','wp:effectExtent'):
        el = inline.find(qn(tag))
        if el is not None: anchor.append(el)
    anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ('wp:docPr','wp:cNvGraphicFramePr'):
        el = inline.find(qn(tag))
        if el is not None: anchor.append(el)
    g = inline.find(qn('a:graphic'))
    if g is not None: anchor.append(g)
    drawing.replace(inline, anchor)


def render_cv(role):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = new_doc(font="Georgia")
    try:
        add_page_bar(doc, os.path.join(_REPO, "documents", "assets", "sidebar_bar.png"), 6.2, 29.7)
    except Exception:
        pass
    for s in doc.sections:
        s.top_margin = Cm(0); s.bottom_margin = Cm(0)
        s.left_margin = Cm(0); s.right_margin = Cm(0)

    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    _no_borders(tbl)
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(21.0 * 567))); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    left, right = tbl.cell(0, 0), tbl.cell(0, 1)
    left.width = Cm(6.2); right.width = Cm(14.8)
    tbl.columns[0].width = Cm(6.2); tbl.columns[1].width = Cm(14.8)
    _shade(left, SIDEBAR_FILL)
    _cell_margins(left, top=150, left=260, bottom=120, right=180)
    _cell_margins(right, top=150, left=300, bottom=120, right=200)
    _row_at_least(tbl.rows[0], 58.0)

    # --- sidebar (never render a heading with nothing under it) ---
    if os.path.exists(HEADSHOT):
        try:
            ph = left.paragraphs[0]; ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ph.paragraph_format.space_after = Pt(6)
            ph.add_run().add_picture(HEADSHOT, width=Cm(3.3))
        except Exception:
            pass
    contact = [x for x in (_PHONE, _EMAIL, _LOCATION, _WEBSITE, _LINKEDIN, *_OTHER_SITES) if x]
    if contact:
        s_head(left, "Contact")
        for line in contact:
            s_text(left, line)
    if SIDEBAR["education"]:
        s_head(left, "Education")
        for line in SIDEBAR["education"]:
            s_text(left, line)
    if SIDEBAR["certificates"]:
        s_head(left, "Certificates")
        for line in SIDEBAR["certificates"]:
            s_text(left, line)
    if SIDEBAR["leadership"]:
        s_head(left, "Leadership")
        for line in SIDEBAR["leadership"]:
            s_text(left, line, bullet=True)
    if SIDEBAR["languages"]:
        s_head(left, "Languages")
        for line in SIDEBAR["languages"]:
            s_text(left, line)
    tags = set(role.get("tags", []))
    if {"ai", "ml", "data"} & tags and SIDEBAR["ai_ml"]:
        s_head(left, "AI / ML")
        for line in SIDEBAR["ai_ml"]:
            s_text(left, line)
    if SIDEBAR["strengths"]:
        s_head(left, "Additional strengths")
        for line in SIDEBAR["strengths"]:
            s_text(left, line, bullet=True)
    tech = next((t for c, t in role["skills"] if c.lower().startswith("tools")), "")
    if tech:
        s_head(left, "Tech")
        s_text(left, tech)

    # --- main ---
    p = right.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(CVD_NAME.upper()); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY; _tracking(r, 20)
    p = right.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(role.get("role_line", role["headline"].split("  ·  ")[0].upper())); r.font.size = Pt(11); r.font.color.rgb = NAVY; _tracking(r, 40)
    p = right.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(role.get("tagline", _CV.get("tagline_default", ""))); r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = GREY

    m_head(right, "Profile")
    m_body(right, role["profile"], size=9.0)
    m_head(right, "Work experience")
    for e in EXPERIENCE:
        extras = [e["extra"][tag] for tag in role["tags"] if tag in e["extra"]]
        m_entry(right, e["org"], e["dates"], "", e["core"], extras, 13.4)
    m_head(right, "Selected work")
    for b in PROJECTS:
        m_bullet(right, b)

    path = os.path.join(OUT, CV_PREFIX + "_CV_" + role["key"] + ".docx")
    doc.save(path); print("wrote", path)

def render_motivation(role):
    doc = new_doc()
    header(doc, role["headline"])
    body(doc, LETTER_DATE, size=9.5, color=GREY, space_after=6)
    recipient = role.get("recipient", "")
    if recipient:
        body(doc, "To: " + recipient, size=9.5, space_after=2)
    body(doc, "Subject: Application for " + role["title"], size=10, bold=True, space_after=6)
    salutation = role.get("salutation", "Dear hiring team,")
    body(doc, salutation, size=9.5, space_after=6)
    body(doc, role["opening"], space_after=6)
    body(doc, role["fit"], space_after=6)
    body(doc, practical_text(role), space_after=6)
    body(doc, "I would welcome a conversation about the role and where I can add the most value. Thank you for your time and consideration.", space_after=10)
    body(doc, "Best regards,", space_after=2)
    body(doc, NAME, bold=True, space_after=0)
    body(doc, MOTIVATION_CONTACT, size=8.5, color=GREY, space_after=0)
    path = os.path.join(OUT, CV_PREFIX + "_Motivation_" + role["key"] + ".docx")
    doc.save(path); print("wrote", path)


def load_archetypes():
    """Role-family templates from config/archetypes.json (or the example)."""
    data = _read_json("archetypes.json")
    return data if isinstance(data, list) else []


ROLES = _CV.get("roles") or load_archetypes()

if __name__ == "__main__":
    import sys
    if "-h" in sys.argv or "--help" in sys.argv:
        print("cv_builder.py is a library. Build documents with tools/build_archetypes.py.")
        sys.exit(0)
    print("cv_builder.py is a library; nothing built. Use tools/build_archetypes.py "
          "(or tools/make_letter.py for a single letter).")
