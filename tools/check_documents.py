#!/usr/bin/env python3
"""Pre-submit gate: refuse to send a defective or AI-sloppy letter.

Run this before every submission. It fails (non-zero exit) if a letter has a
placeholder identity, a missing contact line, a wrong salutation, leftover
template tokens, banned AI-slop wording, or poor ATS keyword coverage.

Usage:
  python3 tools/check_documents.py --letter documents/applications/X.docx \
      --cv documents/archetypes/YOUR_CV.docx \
      --ats-keywords state/ats/<role>.txt

  # loose keyword list instead of a file
  python3 tools/check_documents.py --letter L.docx --keywords "Python,Kubernetes,RAG,MCP"

  # a CV on its own (fixed-CV runs upload no letter): identity, template
  # leftovers, claims about past work and the contract rate
  python3 tools/check_documents.py --cv documents/archetypes/YOUR_CV.docx --record

Exit codes: 0 pass, 1 fail, 2 a document was not found.
"""
import argparse
import os
import re
import sys
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import facts as F
import identity as ID

# Identity comes from config/profile.md (fallback config/profile.example.md).
PROFILE = ID.parse()
NAME = PROFILE.get("full name", "")
CITY = (PROFILE.get("location (city, country)", "").split(",")[0] or "").strip().lower()
CONTRACT_RATE = PROFILE.get("contract rate (ex vat)", "")

# Banned outright by the no-ai-slop skill (word-boundary matched).
SLOP_WORDS = [
    "delve", "foster", "leverage", "utilize", "utilise", "facilitate", "empower",
    "streamline", "robust", "cutting-edge", "paradigm shift", "game changer",
    "game-changer", "this is huge", "this changes everything", "tapestry", "realm",
    "beacon", "multifaceted", "meticulous", "intricate", "paramount", "transformative",
    "elevate", "embark", "supercharge", "ever-evolving",
]
# Empty phrases to avoid.
SLOP_PHRASES = [
    "it's worth noting", "it is worth noting", "it's important to note",
    "at the end of the day", "when it comes to", "at its core", "in today's world",
    "in the age of", "in the world of", "the reality is", "the truth is",
    "with regard to", "going forward", "in conclusion", "let's dive in",
    "here's the thing", "let me be clear", "i'll be honest",
    "underneath sits", "underneath lies", "beneath the surface", "at its foundation",
    "underpinned by", "a long engineering base", "sits a long", "testament to",
]
PLACEHOLDERS = ["your name", "test name", "<company>", "<role>", "<name>",
                "todo", "tbd", "xxx", "lorem", "[insert", "placeholder", "x@y.com",
                "you@example.com", "example.com", "yoursite.com", "linkedin.com/in/you",
                "city, country", "your city", "+00 000"]

# A profile whose identity has not been filled in yet.
PLACEHOLDER_IDENTITIES = {"", "your name", "your_name", "name", "todo", "tbd", "xxx"}

STOPWORDS = set("""a an the and or for with to of in on at as by is are be am was were
we you they their our your this that these those from i my me it its will can could would
have has had not but so if then than about into over under more most other also both each
all any new own no do does did been being""".split())


def docx_text(path):
    doc = Document(path)
    return [p.text for p in doc.paragraphs], doc


def docx_all_paragraphs(path):
    """Every paragraph, including those inside tables: the styled CV is a
    two-column table, so its top-level paragraphs alone are nearly empty."""
    doc = Document(path)
    out = [p.text for p in doc.paragraphs]

    def walk(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    out.extend(p.text for p in cell.paragraphs)
                    walk(cell.tables)
    walk(doc.tables)
    return [t for t in out if t.strip()]


def check_identity(paras, name, errors):
    if not paras:
        errors.append("empty document")
        return
    head = paras[0].strip()
    if not name or name.strip().lower() in PLACEHOLDER_IDENTITIES:
        errors.append("profile identity is a placeholder - set 'full name' in "
                      "config/profile.md (or profile.example.md)")
    elif head.upper() != name.upper():
        errors.append(f"letterhead name is {head!r}, expected {name.upper()!r}")
    body = " ".join(paras)
    low = body.lower()
    for tok in PLACEHOLDERS:
        if tok in low:
            errors.append(f"placeholder token present: {tok!r}")
    contact = paras[2] if len(paras) > 2 else ""
    links = paras[3] if len(paras) > 3 else ""
    if "@" not in contact or not re.search(r"\d{6,}", contact.replace(" ", "")):
        errors.append(f"contact line missing email/phone: {contact!r}")
    if not links.strip():
        errors.append("links line is empty")


def check_cv(cv_paras, name, facts, corpus, errors, warnings):
    """Gate a CV on its own: whose it is, template leftovers, and present-tense
    claims about past work. The letter's prose rules (salutation, slop words,
    repetition, naturalness) are not applied to a CV."""
    low = " ".join(cv_paras).lower()
    if not low.strip():
        errors.append("CV has no text")
        return
    if not name or name.strip().lower() in PLACEHOLDER_IDENTITIES:
        errors.append("profile identity is a placeholder - set 'full name' in "
                      "config/profile.md (or profile.example.md)")
    elif name.lower() not in low:
        errors.append(f"CV does not carry the candidate's name ({name!r})")
    # Whole-word matching: a CV is long enough for "todo" to sit inside "Todoist".
    for tok in PLACEHOLDERS:
        if re.search(r"(?<![a-z0-9])" + re.escape(tok) + r"(?![a-z0-9])", low):
            errors.append(f"placeholder token present: {tok!r}")
    # One paragraph at a time: CV bullets have no full stops, so a joined CV
    # would read as one long sentence mixing current and past employers.
    for para in cv_paras:
        f_err, f_warn = F.check_text(para, facts, corpus)
        errors += f_err
        warnings += f_warn


def check_salutation(paras, expected, company, errors, warnings):
    sals = [p.strip() for p in paras if p.strip().lower().startswith("dear ")]
    if not sals:
        errors.append("no salutation line ('Dear ...')")
        return
    sal = sals[0]
    if expected and expected.lower() not in sal.lower():
        errors.append(f"salutation {sal!r} does not match expected {expected!r}")


def check_slop(paras, errors, warnings):
    body = " ".join(paras)
    low = body.lower()
    for w in SLOP_WORDS:
        if re.search(r"\b" + re.escape(w) + r"\b", body, re.IGNORECASE):
            errors.append(f"banned AI-slop word: {w!r}")
    # "harness" is legitimate as a noun (e.g. "agent harness", "LLM harness");
    # only the slop *verb* usage is banned ("harness the power of ...").
    if re.search(r"\bharness(?:es|ed|ing)?\s+(?:the\s+)?(?:power|potential|benefit|"
                 r"value|opportunit\w*|synerg\w*|strength\w*|advantage\w*)", low) or \
       re.search(r"\bharness(?:ing)?\s+\w+\s+to\s+\w+", low):
        errors.append("banned AI-slop usage: 'harness' as a verb (use it only as a noun, e.g. 'agent harness')")
    for ph in SLOP_PHRASES:
        if ph in low:
            errors.append(f"empty phrase: {ph!r}")
    em = body.count("\u2014")
    dh = body.count(" -- ")
    if em:
        errors.append(f"em dash (\u2014) used {em}x — never use em dashes; use a comma, colon or parentheses")
    if dh:
        errors.append(f"' -- ' used {dh}x — use a comma, colon or parentheses")
    if re.search(r"\bisn't just\b|\bnot just\b|\bnot only\b", low):
        warnings.append("possible binary contrast ('not just X but Y')")


def check_keywords(letter_paras, cv_text, keywords, min_cov, errors, warnings):
    if not keywords:
        return
    corpus = (" ".join(letter_paras) + " " + cv_text).lower()
    missing = [k for k in keywords if k.lower() not in corpus]
    covered = len(keywords) - len(missing)
    cov = covered / len(keywords)
    print(f"  ATS keyword coverage: {covered}/{len(keywords)} ({cov:.0%})")
    if missing:
        print(f"  missing: {', '.join(missing)}")
    if cov < min_cov:
        errors.append(f"ATS keyword coverage {cov:.0%} below threshold {min_cov:.0%}")


def check_claims(letter_paras, facts, corpus, errors, warnings):
    """Fail if a first-person sentence claims scope the CV cannot support.

    The authoritative facts are read from config/cv.json by facts.py; the letter
    is validated with the same generic engine used for free-text answers, so no
    employer, product or domain is hard-coded here.
    """
    f_err, f_warn = F.check_text(" ".join(letter_paras), facts, corpus)
    errors += f_err
    warnings += f_warn


def check_rate(letter_paras, cv_text, engagement, rate, errors):
    """The configured contract rate must not appear in a full-time application."""
    if str(engagement).lower() in ("contract", "contracting", "consulting", "consultant"):
        return
    if not (rate or "").strip():
        return
    text = (" ".join(letter_paras) + " " + cv_text).lower()
    if rate.strip().lower() in text:
        errors.append(f"contract rate ({rate.strip()}) present in a full-time application")


def _body_paragraphs(letter_paras):
    """The prose body: opening, fit and practical paragraphs (no letterhead/subject/signature)."""
    body = []
    for p in letter_paras[4:]:
        t = p.strip()
        if not t:
            continue
        low = t.lower()
        if low.startswith(("subject:", "dear ", "to:", "re:")):
            continue
        if CITY and low.startswith(CITY + ","):
            continue
        if low in ("best regards,", "kind regards,", "sincerely,", "yours sincerely,"):
            continue
        if NAME and (t == NAME or t.upper() == NAME.upper()):
            continue
        if "@" in t or "linkedin.com" in t:
            continue
        body.append(t)
    return body


LOGISTICS = ("available", "immediately", "notice", "based", "citizenship",
             "relocation", "remote", "onsite", "on-site", "commute", "visa", "permit", "located")


def check_duplication(letter_paras, errors, warnings):
    """Repeated logistics and repeated 4+ word phrases are hard fails; short repeats warn."""
    words = re.findall(r"[a-z0-9]+", " ".join(_body_paragraphs(letter_paras)).lower())
    counts = {}
    for n in (2, 3, 4, 5):
        for i in range(len(words) - n + 1):
            gram = words[i:i + n]
            content = [w for w in gram if w not in STOPWORDS and not w.isdigit()]
            if len(content) < max(2, n - 1) or max(len(w) for w in content) < 5:
                continue
            counts[" ".join(gram)] = counts.get(" ".join(gram), 0) + 1
    for phrase, n in counts.items():
        if n < 2:
            continue
        if any(tok in phrase for tok in LOGISTICS):
            errors.append(f"duplicated logistics in letter: '{phrase}' appears {n} times")
        elif len(phrase.split()) >= 4:
            errors.append(f"repeated phrase in letter: '{phrase}' appears {n} times (do not repeat; vary or cut it)")
        else:
            warnings.append(f"repeated phrase in letter: '{phrase}' appears {n} times (consider varying)")


def check_jd_echo(letter_paras, jd_text, errors, n=8):
    """Fail if the letter copies long verbatim runs from the job description."""
    if not jd_text:
        return
    lw = re.findall(r"[a-z0-9]+", " ".join(_body_paragraphs(letter_paras)).lower())
    jw = re.findall(r"[a-z0-9]+", jd_text.lower())
    jgrams = {" ".join(jw[i:i + n]) for i in range(len(jw) - n + 1)}
    hits = []
    for i in range(len(lw) - n + 1):
        g = lw[i:i + n]
        if len([w for w in g if w not in STOPWORDS and not w.isdigit()]) < 4:
            continue
        gram = " ".join(g)
        if gram in jgrams:
            hits.append(gram)
    hits = list(dict.fromkeys(hits))
    if hits:
        errors.append(f"letter echoes the job description verbatim ({len(hits)} phrase(s)); paraphrase in your own words: e.g. '{hits[0]}'")


def check_naturalness(letter_paras, errors, warnings):
    """Heuristics for whether the prose reads naturally (a human would ask this)."""
    body = " ".join(_body_paragraphs(letter_paras))
    sents = [s.strip() for s in re.split(r"(?<=[.!?])\s+", body) if s.strip()]
    long = [s for s in sents if len(s.split()) > 45]
    if long:
        warnings.append(f"{len(long)} very long sentence(s) (>45 words) - may not read naturally")
    starts = {}
    for s in sents:
        parts = s.split()
        if parts:
            starts[parts[0].lower()] = starts.get(parts[0].lower(), 0) + 1
    for w, n in starts.items():
        if n >= 3 and w not in STOPWORDS:
            warnings.append(f"{n} sentences start with '{w}' - vary sentence openings")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--letter", default="", help="cover letter .docx (omit to gate a CV on its own)")
    ap.add_argument("--cv", default="")
    ap.add_argument("--jd", default="", help="job description text file (for reference)")
    ap.add_argument("--ats-keywords", default="", help="file with one keyword per line")
    ap.add_argument("--keywords", default="", help="comma-separated keywords")
    ap.add_argument("--expected-salutation", default="")
    ap.add_argument("--company", default="")
    ap.add_argument("--min-coverage", type=float, default=0.6)
    ap.add_argument("--engagement", default="fulltime", choices=["fulltime", "contract"],
                    help="full-time roles must not carry a contract day rate (default fulltime)")
    ap.add_argument("--answer", action="append", default=[],
                    help="free-text form answer file to validate against the authoritative CV")
    ap.add_argument("--warnings-as-errors", action="store_true")
    ap.add_argument("--record", action="store_true",
                    help="on PASS, write a gate receipt for the CV+letter (hash-stamped); "
                         "record_submission.py requires a matching receipt to record SUBMITTED")
    args = ap.parse_args()

    if not args.letter and not args.cv:
        ap.error("give --letter (with --cv), or --cv alone to gate a CV on its own")
    # A CV that is named but missing must not quietly drop out of the gate.
    for kind, path in (("letter", args.letter), ("CV", args.cv)):
        if path and not os.path.isfile(path):
            print(f"{kind} not found: {path}")
            sys.exit(2)

    errors, warnings = [], []
    cv_paras = docx_all_paragraphs(args.cv) if args.cv else []
    cv_text = " ".join(cv_paras)

    # Authoritative facts: config/cv.json, plus the identity profile as answers.
    facts, corpus = F.load()
    corpus = (corpus + " " + " ".join(ID.parse().values())).lower()

    letter_paras = []
    if args.letter:
        letter_paras, _ = docx_text(args.letter)
        check_identity(letter_paras, NAME, errors)
        check_salutation(letter_paras, args.expected_salutation, args.company, errors, warnings)
        check_slop(letter_paras, errors, warnings)
        check_claims(letter_paras, facts, corpus, errors, warnings)
        check_duplication(letter_paras, errors, warnings)
        check_naturalness(letter_paras, errors, warnings)
        if args.jd and os.path.isfile(args.jd):
            with open(args.jd, encoding="utf-8", errors="ignore") as fh:
                check_jd_echo(letter_paras, fh.read(), errors)
    else:
        check_cv(cv_paras, NAME, facts, corpus, errors, warnings)
    check_rate(letter_paras, cv_text, args.engagement, CONTRACT_RATE, errors)

    if args.answer:
        for path in args.answer:
            if not os.path.isfile(path):
                errors.append(f"answer file not found: {path}")
                continue
            text = open(path, encoding="utf-8").read()
            f_err, f_warn = F.check_text(text, facts, corpus)
            warnings += [f"{os.path.basename(path)}: {w}" for w in f_warn]
            errors += [f"{os.path.basename(path)}: {e}" for e in f_err]

    keywords = []
    if args.keywords:
        keywords = [k.strip() for k in args.keywords.split(",") if k.strip()]
    if args.ats_keywords and os.path.isfile(args.ats_keywords):
        with open(args.ats_keywords) as fh:
            keywords += [l.strip() for l in fh if l.strip() and not l.startswith("#")]
    check_keywords(letter_paras, cv_text, keywords, args.min_coverage, errors, warnings)

    print(f"checked {os.path.basename(args.letter or args.cv)}")
    for w in warnings:
        print(f"  WARN  {w}")
    for e in errors:
        print(f"  FAIL  {e}")
    if errors or (args.warnings_as_errors and warnings):
        print("RESULT: FAIL")
        sys.exit(1)
    if args.record:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import gate_receipt as GR
        rec = GR.emit("documents", {"cv": args.cv, "letter": args.letter},
                      meta={"company": args.company, "engagement": args.engagement})
        if rec:
            print(f"gate receipt: {os.path.relpath(rec, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))}")
        else:
            print("gate receipt: not written (no --cv/--letter files found)")
    print("RESULT: PASS")


if __name__ == "__main__":
    main()
